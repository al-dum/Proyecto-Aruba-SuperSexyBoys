import os
import re
from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = os.path.dirname(os.path.dirname(os.path.realpath(__file__)))
SOURCE = os.path.join(ROOT, "data/external/Table-3.8-Energy-consumption-per-month-in-KwH-by-tariff-group-2015-2024rev.xlsx")
OUTPUT = os.path.join(ROOT, "data/processed/external/energy_consumption_2015_2024_clean.xlsx")

TARIFFS = [
    "Residential",
    "Residential Prepaid",
    "Small Commercial",
    "Large Commercial (hotels)",
    "Streetlight",
    "Security light",
    "Non-commercial",
]
MONTHS = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]
NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
WHITE = "FFFFFF"
GREY = "666666"
THIN_GREY = Side(style="thin", color="B7C9D6")


def style_title(ws, title, subtitle=None):
    ws.merge_cells("A1:I1")
    cell = ws["A1"]
    cell.value = title
    cell.font = Font(name="Arial", size=16, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws.merge_cells("A2:I2")
        cell = ws["A2"]
        cell.value = subtitle
        cell.font = Font(name="Arial", size=10, italic=True, color=GREY)
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[2].height = 20


def style_header(row):
    for cell in row:
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=THIN_GREY, bottom=THIN_GREY)


def add_table(ws, ref, name):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)


def setup_print(ws, print_area, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.print_area = print_area
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.right.text = "N.V. Elmar Aruba"


def annual_rows(ws):
    years = {}
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if isinstance(value, str) and value.startswith("3.8  Energy consumption per month"):
            match = re.search(r"(20\d{2})$", value)
            if match:
                years[int(match.group(1))] = row
    return years


def main():
    source = load_workbook(SOURCE, data_only=True)
    source_ws = source.active
    starts = annual_rows(source_ws)
    if set(starts) != set(range(2015, 2025)):
        raise ValueError("The source workbook does not contain the expected monthly blocks for 2015-2024.")

    monthly = []
    for year, start in starts.items():
        for month_index, month in enumerate(MONTHS, start=1):
            source_row = start + 3 + month_index - 1
            values = [source_ws.cell(source_row, column).value for column in range(2, 9)]
            if source_ws.cell(source_row, 1).value != month or any(value is None for value in values):
                raise ValueError(f"Unexpected monthly data for {year} / {month}.")
            monthly.append([year, month_index, month, *values])

    infrastructure = {}
    for label, title_fragment in (("Usuarios", "Number of users"), ("Conexiones", "Connections by tariff")):
        title_row = next(
            row for row in range(1, source_ws.max_row + 1)
            if isinstance(source_ws.cell(row, 1).value, str) and title_fragment in source_ws.cell(row, 1).value
        )
        infrastructure[label] = [
            [source_ws.cell(row, column).value for column in range(1, 9)]
            for row in range(title_row + 2, title_row + 12)
        ]

    wb = Workbook()
    index = wb.active
    index.title = "Indice"
    style_title(index, "Energy consumption by tariff group", "Cleaned consultation workbook | Period: 2015-2024 | Unit: MWh")
    index["A4"] = "Workbook guide"
    index["A4"].font = Font(name="Arial", size=12, bold=True, color=NAVY)
    index["A5"] = "Use the filtered 'Consumo mensual' sheet for analysis. Each 'Ano YYYY' sheet is a one-page, print-ready view."
    index.merge_cells("A5:I5")
    index["A5"].alignment = Alignment(wrap_text=True)
    index["A7"] = "Sheet"
    index["B7"] = "Purpose"
    style_header(index[7][0:2])
    contents = [
        ("Consumo mensual", "Normalized monthly consumption data; filter by year, month, or tariff group."),
        ("Resumen anual", "Annual totals calculated from the monthly data."),
        ("Usuarios", "Users by tariff group, 2015-2024."),
        ("Conexiones", "Connections by tariff group, 2015-2024."),
    ] + [(f"Ano {year}", f"Monthly consumption for {year}; formatted for one A4 landscape page.") for year in range(2024, 2014, -1)]
    for row, (sheet, description) in enumerate(contents, start=8):
        index.cell(row, 1, sheet).hyperlink = f"#'{sheet}'!A1"
        index.cell(row, 1).style = "Hyperlink"
        index.cell(row, 2, description)
        index.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    index["A23"] = "Source"
    index["B23"] = "N.V. Elmar Aruba. The source labels its values as MWh; this workbook retains that unit."
    index["A23"].font = Font(name="Arial", bold=True)
    index["B23"].alignment = Alignment(wrap_text=True)
    index.column_dimensions["A"].width = 24
    index.column_dimensions["B"].width = 90
    index.sheet_view.showGridLines = False
    index.freeze_panes = "A7"

    data = wb.create_sheet("Consumo mensual")
    style_title(data, "Monthly energy consumption", "One record per year and month | Unit: MWh | Filter the table to consult or clean the data.")
    headers = ["Year", "Month number", "Month", *TARIFFS, "Total"]
    for col, header in enumerate(headers, start=1):
        data.cell(4, col, header)
    style_header(data[4])
    for row_number, row in enumerate(sorted(monthly), start=5):
        for column, value in enumerate(row, start=1):
            data.cell(row_number, column, value)
        data.cell(row_number, 11, f"=SUM(D{row_number}:J{row_number})")
    for row in data.iter_rows(min_row=5, max_row=124, min_col=4, max_col=11):
        for cell in row:
            cell.number_format = '#,##0.00;[Red](#,##0.00);-'
    widths = [10, 14, 14, 16, 21, 19, 27, 14, 16, 18, 16]
    for column, width in enumerate(widths, start=1):
        data.column_dimensions[get_column_letter(column)].width = width
    add_table(data, "A4:K124", "MonthlyConsumption")
    data.conditional_formatting.add("K5:K124", ColorScaleRule(start_type="min", start_color="FFFFFF", mid_type="percentile", mid_value=50, mid_color="9DC3E6", end_type="max", end_color="1F4E78"))
    data.sheet_view.showGridLines = False
    data.freeze_panes = "A5"
    data.auto_filter.ref = "A4:K124"

    summary = wb.create_sheet("Resumen anual")
    style_title(summary, "Annual energy consumption", "Calculated from the monthly consumption table | Unit: MWh")
    annual_headers = ["Year", *TARIFFS, "Total"]
    for col, header in enumerate(annual_headers, start=1):
        summary.cell(4, col, header)
    style_header(summary[4])
    for output_row, year in enumerate(range(2015, 2025), start=5):
        summary.cell(output_row, 1, year)
        for output_col in range(2, 10):
            source_col = get_column_letter(output_col + 2)
            summary.cell(output_row, output_col, f'=SUMIF(\'Consumo mensual\'!$A$5:$A$124,A{output_row},\'Consumo mensual\'!${source_col}$5:${source_col}$124)')
        summary.cell(output_row, 9, f"=SUM(B{output_row}:H{output_row})")
    for row in summary.iter_rows(min_row=5, max_row=14, min_col=2, max_col=9):
        for cell in row:
            cell.number_format = '#,##0.00;[Red](#,##0.00);-'
    for column, width in enumerate([12, 16, 21, 19, 27, 14, 16, 18, 16], start=1):
        summary.column_dimensions[get_column_letter(column)].width = width
    add_table(summary, "A4:I14", "AnnualConsumption")
    setup_print(summary, "A1:I14")

    for sheet_name, values in infrastructure.items():
        ws = wb.create_sheet(sheet_name)
        style_title(ws, f"{sheet_name} by tariff group", "Period: 2015-2024 | Counts")
        for col, header in enumerate(["Year", *TARIFFS, "Total"], start=1):
            ws.cell(4, col, header)
        style_header(ws[4])
        for output_row, row in enumerate(values, start=5):
            for output_col, value in enumerate(row, start=1):
                ws.cell(output_row, output_col, value)
            ws.cell(output_row, 9, f"=SUM(B{output_row}:H{output_row})")
        for row in ws.iter_rows(min_row=5, max_row=14, min_col=2, max_col=9):
            for cell in row:
                cell.number_format = '#,##0;[Red](#,##0);-'
        for column, width in enumerate([12, 16, 21, 19, 27, 14, 16, 18, 16], start=1):
            ws.column_dimensions[get_column_letter(column)].width = width
        add_table(ws, "A4:I14", sheet_name.replace(" ", ""))
        setup_print(ws, "A1:I14")

    for year in range(2024, 2014, -1):
        ws = wb.create_sheet(f"Ano {year}")
        style_title(ws, f"Energy consumption by tariff group - {year}", "Monthly detail | Unit: MWh")
        for col, header in enumerate(["Month", *TARIFFS, "Total"], start=1):
            ws.cell(4, col, header)
        style_header(ws[4])
        year_rows = [row for row in monthly if row[0] == year]
        for output_row, row in enumerate(year_rows, start=5):
            ws.cell(output_row, 1, row[2])
            for output_col, value in enumerate(row[3:], start=2):
                ws.cell(output_row, output_col, value)
            ws.cell(output_row, 9, f"=SUM(B{output_row}:H{output_row})")
        total_row = 17
        ws.cell(total_row, 1, "Annual total")
        for column in range(2, 10):
            letter = get_column_letter(column)
            ws.cell(total_row, column, f"=SUM({letter}5:{letter}16)")
        for cell in ws[total_row]:
            cell.font = Font(name="Arial", bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.border = Border(top=THIN_GREY, bottom=THIN_GREY)
        for row in ws.iter_rows(min_row=5, max_row=17, min_col=2, max_col=9):
            for cell in row:
                cell.number_format = '#,##0.00;[Red](#,##0.00);-'
        for column, width in enumerate([15, 16, 21, 19, 27, 14, 16, 18, 16], start=1):
            ws.column_dimensions[get_column_letter(column)].width = width
        add_table(ws, "A4:I16", f"Year{year}")
        setup_print(ws, "A1:I17")

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name != "Arial":
                    font = copy(cell.font)
                    font.name = "Arial"
                    cell.font = font

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
